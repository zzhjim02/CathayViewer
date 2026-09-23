# -*- coding: utf-8 -*-
"""CathayViewer · 学术书库浏览与阅读 —— 人名别名归一  v0.1.0

近代人物常有多个 字 / 号 / 笔名 / 化名 / 绰号。本模块做「别名归一」：
  * lookup(kw)  → 这个关键词是哪个（些）人，别名有哪些
  * expand(kw)  → 检索时该一并搜的其它名字（如 吴佩孚 → 子玉、吴玉帅…）
  * import_csv/export_csv → 离线扩充（把 CBDB / Wikidata / 中研院人名档 的导出并进来）

纯逻辑、不依赖 PyQt、不联网。数据 = 内置精简表 + 用户本地的 person_alias.csv。
"""
import os

# 正名 → [字/号/笔名/化名/绰号...]（可核验的公开事实；用户可自行增补）
DEFAULT_ALIASES = {
    # —— 先秦—汉 ——
    '孔子': ['孔丘', '仲尼', '孔夫子', '孔圣人', '至圣先师', '素王', '尼父'],
    '孟子': ['孟轲', '子舆', '亚圣', '孟夫子'],
    '老子': ['李耳', '老聃', '李聃', '太上老君', '道德天尊'],
    '庄子': ['庄周', '子休', '南华真人'],
    '荀子': ['荀况', '荀卿', '孙卿'],
    '屈原': ['屈平', '灵均', '三闾大夫'],
    '司马迁': ['子长', '太史公', '史迁'],
    '曹操': ['孟德', '曹孟德', '魏武帝', '阿瞒', '曹阿瞒'],
    '刘备': ['玄德', '刘玄德', '昭烈帝', '先主'],
    '孙权': ['仲谋', '孙仲谋'],
    '诸葛亮': ['孔明', '诸葛孔明', '卧龙', '卧龙先生', '忠武侯', '武乡侯'],
    '关羽': ['云长', '关云长', '关公', '关帝', '武圣', '美髯公', '汉寿亭侯'],
    '张飞': ['翼德', '张翼德'],
    '周瑜': ['公瑾', '周郎', '周公瑾'],
    # —— 唐—清 ——
    '陶渊明': ['陶潜', '元亮', '陶元亮', '五柳先生', '靖节先生'],
    '李白': ['李太白', '太白', '青莲居士', '诗仙', '谪仙人', '李十二'],
    '杜甫': ['杜子美', '子美', '少陵野老', '杜工部', '杜拾遗', '诗圣', '杜少陵'],
    '白居易': ['白乐天', '乐天', '香山居士', '醉吟先生', '白香山'],
    '王维': ['王摩诘', '摩诘', '诗佛'],
    '韩愈': ['韩退之', '退之', '昌黎', '韩昌黎', '韩文公', '昌黎先生'],
    '柳宗元': ['柳子厚', '子厚', '柳河东', '柳柳州'],
    '苏轼': ['苏东坡', '苏子瞻', '子瞻', '东坡居士', '和仲', '铁冠道人'],
    '欧阳修': ['欧阳永叔', '永叔', '醉翁', '六一居士'],
    '王安石': ['王介甫', '介甫', '半山', '临川先生', '荆公'],
    '陆游': ['陆放翁', '务观', '放翁'],
    '辛弃疾': ['辛稼轩', '幼安', '稼轩'],
    '岳飞': ['岳鹏举', '鹏举', '武穆', '岳武穆'],
    '文天祥': ['文文山', '文山', '宋瑞', '履善', '信国公', '文信国公'],
    '朱熹': ['朱子', '朱文公', '元晦', '仲晦', '晦庵', '晦翁', '紫阳', '考亭'],
    '包拯': ['包公', '包孝肃', '包青天', '包希仁'],
    '范仲淹': ['范文正', '希文', '文正公'],
    '王阳明': ['王守仁', '守仁', '伯安', '阳明', '阳明先生'],
    '曹雪芹': ['曹霑', '梦阮', '雪芹', '芹圃', '芹溪'],
    '蒲松龄': ['留仙', '剑臣', '柳泉居士', '聊斋先生'],
    '吴敬梓': ['敏轩', '粒民', '文木老人'],
    # —— 晚清—民国（近代）——
    '李鸿章': ['李少荃', '少荃', '渐甫', '李中堂', '李合肥', '仪叟'],
    '曾国藩': ['曾伯涵', '伯涵', '涤生', '曾文正', '曾涤生'],
    '左宗棠': ['左季高', '季高', '湘上农人', '老亮'],
    '张之洞': ['张香涛', '香涛', '孝达', '香岩', '无竞居士', '张南皮'],
    '严复': ['严几道', '几道', '又陵', '尊疑', '尺庵'],
    '康有为': ['康南海', '康长素', '广厦', '长素', '南海', '明夷', '西樵山人', '天游化人'],
    '梁启超': ['梁任公', '梁卓如', '卓如', '任甫', '任公', '饮冰室主人', '饮冰子', '哀时客'],
    '谭嗣同': ['谭壮飞', '复生', '壮飞', '华相众生'],
    '孙中山': ['孙文', '孙逸仙', '逸仙', '载之', '中山樵', '帝象', '德明', '孙大炮'],
    '黄兴': ['黄克强', '克强', '黄轸', '庆午', '堇午'],
    '蔡锷': ['蔡松坡', '松坡', '艮寅', '奋翮生', '击椎生'],
    '袁世凯': ['袁慰亭', '慰亭', '慰廷', '容庵', '袁项城', '项城', '袁宫保', '洹上钓叟'],
    '黎元洪': ['黎宋卿', '宋卿'],
    '段祺瑞': ['段芝泉', '芝泉', '正道老人'],
    '冯国璋': ['冯华甫', '华甫'],
    '徐世昌': ['徐东海', '卜五', '水竹邨人', '菊人'],
    '曹锟': ['曹仲珊', '仲珊'],
    '张作霖': ['张雨亭', '雨亭', '老帅', '张大帅', '雨帅', '东北王'],
    '张学良': ['张汉卿', '汉卿', '少帅', '毅庵'],
    '吴佩孚': ['子玉', '吴玉帅', '玉帅', '孚威将军', '孚威'],
    '孙传芳': ['孙馨远', '馨远', '五省联军总司令'],
    '阎锡山': ['阎百川', '百川', '龙池'],
    '冯玉祥': ['冯焕章', '焕章', '基督将军', '倒戈将军'],
    '李宗仁': ['李德邻', '德邻'],
    '白崇禧': ['白健生', '健生', '小诸葛'],
    '傅作义': ['宜生'],
    '蒋介石': ['蒋中正', '中正', '介石', '瑞元', '志清', '蒋校长', '蒋委员长', '蒋公'],
    '汪精卫': ['汪兆铭', '兆铭', '季新', '精卫', '汪季新'],
    '陈炯明': ['陈竞存', '竞存'],
    '廖仲恺': ['廖恩煦', '恩煦', '夷白'],
    '宋教仁': ['宋遁初', '遁初', '渔父'],
    '蔡元培': ['蔡孑民', '鹤卿', '仲申', '民友', '孑民', '鹤庼', '周子余'],
    '章太炎': ['章炳麟', '炳麟', '枚叔', '太炎', '菿汉', '余杭先生'],
    '王国维': ['王静安', '静安', '伯隅', '观堂', '静庵', '永观', '王观堂'],
    '陈寅恪': ['鹤寿'],
    '辜鸿铭': ['辜汤生', '汤生', '鸿铭', '汉滨读易者'],
    '鲁迅': ['周树人', '周豫才', '豫才', '豫山', '庚辰', '自树', '唐俟', '巴人',
             '何家干', '乐雯', '洛文', '游光', '葛何德', '封余', '俟堂'],
    '胡适': ['胡洪骍', '洪骍', '嗣穈', '适之', '期自胜生', '铁儿', '藏晖', '天风'],
    '陈独秀': ['陈仲甫', '仲甫', '独秀', '乾生', '由己', '三爱', '只眼', '撒翁'],
    '李大钊': ['李守常', '守常', '龟年', '耆年'],
    '林纾': ['林琴南', '琴南', '畏庐', '冷红生', '践卓翁'],
    '熊十力': ['继智', '子真', '漆园老人', '逸翁'],
    '马一浮': ['马浮', '一浮', '湛翁', '蠲叟', '蠲戏老人'],
    '李叔同': ['弘一', '弘一法师', '叔同', '息霜', '演音', '晚晴老人'],
    # —— 清帝后（常用称谓）——
    '康熙': ['玄烨', '爱新觉罗·玄烨', '清圣祖'],
    '乾隆': ['弘历', '爱新觉罗·弘历', '清高宗'],
    '光绪': ['载湉', '爱新觉罗·载湉', '清德宗'],
    '慈禧': ['叶赫那拉', '慈禧太后', '西太后', '老佛爷', '那拉氏'],
    '溥仪': ['宣统', '宣统帝', '康德皇帝', '末代皇帝'],
}

_SEPS = ',，、;；/|'


def _norm(s):
    return ''.join(str(s or '').split()).strip()


def _split_aliases(s):
    out, cur = [], ''
    for ch in str(s or ''):
        if ch in _SEPS:
            if cur.strip():
                out.append(_norm(cur))
            cur = ''
        else:
            cur += ch
    if cur.strip():
        out.append(_norm(cur))
    return [x for x in out if x]


# ----------------------------------------------------------------- 载入 / 合并
_groups = None            # 正名 -> [别名]
_idx = None               # 别名/正名 -> {正名...}
_sources = []             # 载入过的文件（供界面显示）


def _app_dir():
    try:
        import viewer_core as C
        return C.app_dir()
    except Exception:
        return os.path.dirname(os.path.abspath(__file__))


def user_csv_path():
    return os.path.join(_app_dir(), 'person_alias.csv')


def bundled_csv_path():
    return os.path.join(_app_dir(), 'config', 'person_alias.csv')


def _read_csv(path):
    """读 CSV → {正名:[别名...]}；逗号/顿号/分号/斜杠/竖线都可作分隔，'#' 起始为注释。"""
    out = {}
    try:
        with open(path, 'r', encoding='utf-8-sig') as f:
            lines = f.read().splitlines()
    except Exception:
        return out
    for ln in lines:
        ln = ln.strip()
        if not ln or ln.startswith('#'):
            continue
        parts = _split_aliases(ln)
        if len(parts) < 2:
            continue
        canon = parts[0]
        out.setdefault(canon, [])
        for a in parts[1:]:
            if a != canon and a not in out[canon]:
                out[canon].append(a)
    return out


def reload(extra_paths=None):
    """重建别名表：内置 + 用户 person_alias.csv + 可选外部 CSV。"""
    global _groups, _sources, _idx
    g = {k: list(v) for k, v in DEFAULT_ALIASES.items()}
    _sources = []
    _idx = None
    paths = []
    bp = bundled_csv_path()
    if os.path.isfile(bp):
        paths.append(bp)
    paths.append(user_csv_path())
    if extra_paths:
        paths.extend(extra_paths)
    for p in paths:
        d = _read_csv(p)
        if d:
            _sources.append(p)
        for canon, als in d.items():
            g.setdefault(canon, [])
            for a in als:
                if a != canon and a not in g[canon]:
                    g[canon].append(a)
    _groups = g
    return g


def groups():
    if _groups is None:
        reload()
    return _groups


def sources():
    return list(_sources)


def count():
    g = groups()
    return {'people': len(g), 'aliases': sum(len(v) for v in g.values())}


def _index():
    global _idx
    if _idx is not None:
        return _idx
    idx = {}
    for canon, als in groups().items():
        idx.setdefault(canon, set()).add(canon)
        for a in als:
            idx.setdefault(a, set()).add(canon)
    _idx = idx
    return idx


# ----------------------------------------------------------------- 查询
def lookup(kw):
    """关键词 → [{'canonical','aliases','matched_alias'}]（精确匹配正名或别名）。"""
    kw = _norm(kw)
    if not kw:
        return []
    idx = _index()
    g = groups()
    out = []
    for canon in sorted(idx.get(kw, ())):
        als = [a for a in g.get(canon, []) if a != kw]
        out.append({'canonical': canon, 'aliases': als,
                    'matched_alias': '' if canon == kw else kw})
    return out


def expand(kw, limit=16):
    """要一并检索的其它名字（保序去重，不含 kw 本身）。引用了别名时也含正名。

    多个人同名号时汇总，最多 limit 个（避免 CBDB 里大量同名号把结果筛爆）。
    """
    out = []
    for it in lookup(kw):
        cn = it['canonical']
        if cn != kw and cn not in out:
            out.append(cn)
        for a in it['aliases']:
            if a != kw and a not in out:
                out.append(a)
        if len(out) >= limit:
            break
    return out[:limit]


def canonical_of(kw):
    hits = lookup(kw)
    return hits[0]['canonical'] if hits else ''


# ----------------------------------------------------------------- 维护
def write_csv(path, data=None):
    """把别名表写成 CSV（正名,别名1,别名2,...）。"""
    data = data or groups()
    lines = ['# CathayViewer 人名别名表（正名,别名…；可手动编辑，逗号/顿号/分号均可分隔）']
    for canon in sorted(data.keys()):
        lines.append(canon + ',' + ','.join(data[canon]))
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    return path


def import_csv(path, merge=True):
    """离线并入一个 CSV 到用户表；返回 (新增人数, 新增别名数)。"""
    d = _read_csv(path)
    if not d:
        return (0, 0)
    cur = {}
    up = user_csv_path()
    if os.path.isfile(up):
        cur = _read_csv(up)
    newp = newa = 0
    for canon, als in d.items():
        cur.setdefault(canon, [])
        for a in als:
            if a != canon and a not in cur[canon]:
                cur[canon].append(a)
                newa += 1
        if not cur[canon]:
            newp += 1
    if merge:
        newp = len([c for c in d if c not in _read_csv(up)]) if os.path.isfile(up) else len(d)
    write_csv(up, cur)
    reload()
    return (newp, newa)


def add_person(canon, aliases):
    """增/改一个人（写入用户表）。返回该人别名数。"""
    canon = _norm(canon)
    if not canon:
        return 0
    als = [a for a in (_split_aliases(aliases) if isinstance(aliases, str) else aliases)
           if a and a != canon]
    up = user_csv_path()
    cur = _read_csv(up) if os.path.isfile(up) else {}
    cur[canon] = als
    write_csv(up, cur)
    reload()
    return len(als)


def remove_person(canon):
    up = user_csv_path()
    cur = _read_csv(up) if os.path.isfile(up) else {}
    if canon in cur:
        del cur[canon]
        write_csv(up, cur)
        reload()
        return True
    return False


# ----------------------------------------------------------------- CBDB 离线导入
def _is_cjk_name(s):
    if not s:
        return False
    cjk = sum(1 for ch in s if '\u4e00' <= ch <= '\u9fff')
    return cjk >= 2 and cjk >= len(s) * 0.6


def import_cbdb_dir(folder, merge=True, min_len=2, max_aliases=40,
                    year_from=None, year_to=None, on_progress=None):
    """从 CBDB 导出的表目录（含 ALTNAME_DATA.xlsx / BIOG_MAIN.xlsx）生成别名表。

    仅读用户自己的 CBDB 导出，产出写入本地 person_alias.csv（不随程序发布）。
    返回 (人数, 别名数, 输出路径)。需要 openpyxl。
    """
    try:
        from openpyxl import load_workbook
    except Exception:
        raise RuntimeError('需要 openpyxl：请先 pip install openpyxl')

    def _find(pref):
        try:
            for fn in os.listdir(folder):
                fl = fn.lower()
                if fl.startswith(pref.lower()) and fl.endswith('.xlsx'):
                    return os.path.join(folder, fn)
        except OSError:
            pass
        return ''

    f_alt = _find('ALTNAME_DATA')
    f_bio = _find('BIOG_MAIN')
    if not f_alt or not f_bio:
        raise RuntimeError('目录里找不到 ALTNAME_DATA.xlsx / BIOG_MAIN.xlsx')

    def _head(ws):
        it = ws.iter_rows(values_only=True)
        h = [('' if c is None else str(c)) for c in next(it)]
        return it, h

    def _ci(h, name):
        try:
            return h.index(name)
        except ValueError:
            return -1

    # 1) BIOG_MAIN：personid → 正名（兼记索引年）
    canon, yb, n = {}, {}, 0
    wb = load_workbook(f_bio, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it, h = _head(ws)
    i_pid, i_chn, i_yr = _ci(h, 'c_personid'), _ci(h, 'c_name_chn'), _ci(h, 'c_index_year')
    for row in it:
        pid = row[i_pid] if i_pid >= 0 else None
        chn = row[i_chn] if i_chn >= 0 else None
        if pid is None or not chn:
            continue
        canon[pid] = str(chn).strip()
        if i_yr >= 0:
            try:
                yb[pid] = int(row[i_yr])
            except Exception:
                pass
        n += 1
        if on_progress and n % 20000 == 0:
            on_progress(n, 0)
    wb.close()

    # 2) ALTNAME_DATA：personid → [别名]
    groups, m = {}, 0
    wb = load_workbook(f_alt, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it, h = _head(ws)
    i_pid, i_chn = _ci(h, 'c_personid'), _ci(h, 'c_alt_name_chn')
    for row in it:
        pid = row[i_pid] if i_pid >= 0 else None
        chn = row[i_chn] if i_chn >= 0 else None
        if pid is None or not chn:
            continue
        if year_from or year_to:
            y = yb.get(pid)
            if y is None or (year_from and y < year_from) or (year_to and y > year_to):
                continue
        c = canon.get(pid)
        s = str(chn).strip()
        if not c or not _is_cjk_name(s) or s == c or len(s) < min_len:
            continue
        lst = groups.setdefault(pid, [])
        if s not in lst and len(lst) < max_aliases:
            lst.append(s)
        m += 1
        if on_progress and m % 50000 == 0:
            on_progress(n, m)
    wb.close()

    data = {}
    for pid, als in groups.items():
        c = canon.get(pid)
        if c and als:
            data[c] = als
    out = user_csv_path()
    if merge and os.path.isfile(out):
        cur = _read_csv(out)
        for c, als in data.items():
            cur.setdefault(c, [])
            for a in als:
                if a != c and a not in cur[c]:
                    cur[c].append(a)
        data = cur
    write_csv(out, data)
    reload()
    return (len(data), sum(len(v) for v in data.values()), out)


# ----------------------------------------------------------------- 自检
def selftest():
    import shutil
    import sys
    import tempfile
    log = []

    def ok(c, m):
        log.append(('OK  ' if c else 'FAIL') + ' ' + m)
        return bool(c)

    base = tempfile.mkdtemp(prefix='va_')
    old = _app_dir
    try:
        globals()['_app_dir'] = lambda: base        # 用户表落到临时目录
        reload()
        c = count()
        ok(c['people'] >= 60 and c['aliases'] >= 200,
           '内置别名表：%d 人 / %d 别名' % (c['people'], c['aliases']))
        wpf = expand('吴佩孚')
        ok(all(x in wpf for x in ('子玉', '吴玉帅', '玉帅', '孚威将军', '孚威')),
           '吴佩孚 → %s' % wpf)
        szs = expand('孙中山')
        ok(all(x in szs for x in ('孙文', '逸仙', '中山樵')), '孙中山 → %s' % szs)
        ok(canonical_of('子玉') == '吴佩孚' and canonical_of('孙文') == '孙中山',
           '反查 子玉/孙文 → 吴佩孚/孙中山')
        ok(expand('周树人') and '鲁迅' in expand('周树人'), '鲁迅↔周树人')
        ok(expand('不存在的名字zz') == [], '未收录名字 → 空')
        # 增 / 改 / 删 / 导入导出
        add_person('测试人物', '甲一,乙二、丙三；丁四')
        ok(set(expand('测试人物')) == {'甲一', '乙二', '丙三', '丁四'},
           '新增人名（多种分隔符）：%s' % expand('测试人物'))
        ex = os.path.join(base, 'exp.csv')
        write_csv(ex)
        ok(os.path.isfile(ex) and '吴佩孚' in open(ex, encoding='utf-8').read(),
           '导出 CSV')
        imp = os.path.join(base, 'imp.csv')
        open(imp, 'w', encoding='utf-8').write('导入人物,别名甲,别名乙\n')
        n = import_csv(imp)
        ok(canonical_of('别名甲') == '导入人物' and n[1] >= 2, '导入 CSV：%s' % (n,))
        remove_person('测试人物')
        ok(expand('测试人物') == [], '删除人名')
    finally:
        globals()['_app_dir'] = old
        shutil.rmtree(base, ignore_errors=True)
    out = '\n'.join(log) + '\nresult = %s\n' % (
        'OK' if all(l.startswith('OK') for l in log) else 'FAIL')
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    print('CathayViewer · 人名别名归一  v0.1.0')
    print(out)


if __name__ == '__main__':
    selftest()
